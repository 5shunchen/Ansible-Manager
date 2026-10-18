"""Agent 主机作用域 + 主机导入导出测试（Excel）。"""
import io

import pytest
from openpyxl import Workbook, load_workbook

HEADERS = ["主机名", "IP地址", "SSH端口", "主机类型", "登录凭据", "所属项目", "描述"]


def _xlsx_bytes(rows: list[list]) -> bytes:
    buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(buf)
    return buf.getvalue()


def _read_xlsx(content: bytes) -> list[list]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


async def _credential(client, auth_headers, name: str) -> int:
    resp = await client.post("/api/v1/credentials", headers=auth_headers, json={
        "name": name, "auth_type": "password", "username": "root", "password": "secret",
    })
    assert resp.status_code == 200, resp.text
    return resp.json()["data"]["id"]


def _scoped(auth_headers, scope):
    headers = dict(auth_headers)
    if scope is not None:
        headers["X-Agent-Host-Id"] = str(scope)
    return headers


async def _host(client, auth_headers, *, hostname, ip, cred_id, host_type="managed", scope=None):
    resp = await client.post("/api/v1/hosts", headers=_scoped(auth_headers, scope), json={
        "hostname": hostname, "ip": ip, "host_type": host_type, "credential_id": cred_id,
    })
    assert resp.status_code == 200, resp.text
    return resp.json()["data"]["id"]


async def _hosts(client, auth_headers, scope=None, **params):
    resp = await client.get(
        "/api/v1/hosts", headers=_scoped(auth_headers, scope),
        params={"page": 1, "page_size": 50, **params},
    )
    assert resp.status_code == 200, resp.text
    return {h["hostname"] for h in resp.json()["data"]["items"]}


async def test_scope_narrows_hosts_and_auto_assigns(client, auth_headers):
    """列表按当前 Agent 主机收窄；新建主机自动归属到该 Agent 主机。"""
    cred = await _credential(client, auth_headers, "c-scope")
    a = await _host(client, auth_headers, hostname="agent-a", ip="10.1.0.1", cred_id=cred, host_type="agent")
    b = await _host(client, auth_headers, hostname="agent-b", ip="10.1.0.2", cred_id=cred, host_type="agent")
    await _host(client, auth_headers, hostname="under-a", ip="10.2.0.1", cred_id=cred, scope=a)
    await _host(client, auth_headers, hostname="under-b", ip="10.2.0.2", cred_id=cred, scope=b)

    # 未选 Agent 主机：不隐藏数据（保证升级期可用）
    assert await _hosts(client, auth_headers) >= {"agent-a", "agent-b", "under-a", "under-b"}

    # 选中 A：只看 A 名下的被管主机，Agent 主机自身不混入
    assert await _hosts(client, auth_headers, scope=a) == {"under-a"}
    assert await _hosts(client, auth_headers, scope=b) == {"under-b"}


async def test_agent_host_list_is_not_narrowed(client, auth_headers):
    """Agent 主机列表本身是作用域维度，选择后仍能看到全部，否则无法切换。"""
    cred = await _credential(client, auth_headers, "c-agent")
    a = await _host(client, auth_headers, hostname="ag-1", ip="10.5.0.1", cred_id=cred, host_type="agent")
    await _host(client, auth_headers, hostname="ag-2", ip="10.5.0.2", cred_id=cred, host_type="agent")

    names = await _hosts(client, auth_headers, scope=a, host_type="agent")
    assert {"ag-1", "ag-2"} <= names


async def test_scope_narrows_projects(client, auth_headers):
    """项目同样按 Agent 主机收窄。"""
    cred = await _credential(client, auth_headers, "c-proj")
    a = await _host(client, auth_headers, hostname="ag-p", ip="10.6.0.1", cred_id=cred, host_type="agent")

    resp = await client.post("/api/v1/projects", headers=_scoped(auth_headers, a), json={"name": "proj-a", "code": "proj-a"})
    assert resp.status_code == 200, resp.text
    resp = await client.post("/api/v1/projects", headers=auth_headers, json={"name": "proj-global", "code": "proj-global"})
    assert resp.status_code == 200, resp.text

    resp = await client.get("/api/v1/projects", headers=_scoped(auth_headers, a))
    assert resp.status_code == 200, resp.text
    names = {p["name"] for p in resp.json()["data"]}
    assert names == {"proj-a"}


async def test_first_agent_host_takes_over_history(client, auth_headers):
    """升级前的无归属数据，在建成第一台 Agent 主机时自动归到它名下。"""
    cred = await _credential(client, auth_headers, "c-backfill")
    # 模拟升级前的历史数据：建主机时不带作用域，归属为空
    await _host(client, auth_headers, hostname="legacy-1", ip="10.7.0.1", cred_id=cred)
    project = await client.post("/api/v1/projects", headers=auth_headers, json={"name": "legacy-proj", "code": "legacy-proj"})
    assert project.status_code == 200, project.text

    a = await _host(client, auth_headers, hostname="ag-first", ip="10.7.0.9", cred_id=cred, host_type="agent")

    # 归到新 Agent 主机后，选中它就能看到历史数据
    assert await _hosts(client, auth_headers, scope=a) == {"legacy-1"}
    resp = await client.get("/api/v1/projects", headers=_scoped(auth_headers, a))
    assert {p["name"] for p in resp.json()["data"]} == {"legacy-proj"}


async def test_import_template_and_export(client, auth_headers):
    """模板与导出同表头，导出内容可直接回灌导入。"""
    cred = await _credential(client, auth_headers, "c-io")
    await _host(client, auth_headers, hostname="io-1", ip="10.3.0.1", cred_id=cred)

    resp = await client.get("/api/v1/hosts/import-template", headers=auth_headers)
    assert resp.status_code == 200
    assert "attachment" in resp.headers["content-disposition"]
    rows = _read_xlsx(resp.content)
    assert list(rows[0]) == HEADERS
    assert len(rows) == 2  # 表头 + 示例行

    resp = await client.get("/api/v1/hosts/export", headers=auth_headers)
    assert resp.status_code == 200
    rows = _read_xlsx(resp.content)
    assert list(rows[0]) == HEADERS
    assert any(str(r[0]) == "io-1" and str(r[1]) == "10.3.0.1" and "c-io" in str(r[4]) for r in rows[1:])


async def test_import_reports_per_row_failures(client, auth_headers):
    """逐行导入：成功行落库，失败行给出行号与原因，且归属当前作用域。"""
    cred = await _credential(client, auth_headers, "c-imp")
    a = await _host(client, auth_headers, hostname="ag-imp", ip="10.4.0.1", cred_id=cred, host_type="agent")
    await _host(client, auth_headers, hostname="io-dup-src", ip="10.4.0.9", cred_id=cred)

    body = _xlsx_bytes([
        ["io-new", "10.4.0.2", "22", "managed", "c-imp", "", "批量导入"],
        ["io-dup", "10.4.0.9", "22", "managed", "c-imp", "", "IP 重复"],
        ["io-bad", "10.4.0.3", "22", "managed", "nope", "", "凭据不存在"],
    ])
    resp = await client.post(
        "/api/v1/hosts/import", headers=_scoped(auth_headers, a),
        files={"file": ("hosts.xlsx", body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()["data"]
    assert (data["total"], data["success"], data["failed"]) == (3, 1, 2)
    assert [e["line"] for e in data["errors"]] == [3, 4]
    assert "已存在" in data["errors"][0]["message"]
    assert "凭据不存在" in data["errors"][1]["message"]

    assert await _hosts(client, auth_headers, scope=a) == {"io-new"}


async def test_import_rejects_wrong_header(client, auth_headers):
    """表头不符合模板时整体拒绝，避免建出脏数据。"""
    buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.append(["名字", "地址"])
    ws.append(["x", "1.1.1.1"])
    wb.save(buf)
    resp = await client.post(
        "/api/v1/hosts/import", headers=auth_headers,
        files={"file": ("bad.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code >= 400
    assert "表头缺少列" in resp.json()["message"]


async def test_import_requires_permission(client, auth_headers):
    """无 host:create 权限的只读用户不能导入。"""
    resp = await client.get("/api/v1/roles", headers=auth_headers)
    viewer = next(r for r in resp.json()["data"] if r["code"] == "viewer")
    await client.post("/api/v1/users", headers=auth_headers, json={
        "username": "imp-viewer", "password": "viewer123", "role_ids": [viewer["id"]],
    })
    login = await client.post("/api/v1/auth/login", json={"username": "imp-viewer", "password": "viewer123"})
    viewer_headers = {"Authorization": f"Bearer {login.json()['data']['token']}"}

    resp = await client.post(
        "/api/v1/hosts/import", headers=viewer_headers,
        files={"file": ("hosts.xlsx", _xlsx_bytes([]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 403
